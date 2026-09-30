#!/usr/bin/env python3
"""Export the corporation valuation to a spreadsheet (corp_values.xlsx) + CSV.

Sheets:
  Corp values         one row per corp: deltas, observed MC value, tangible,
                      implied ability value, description
  Tangible breakdown  component counts and their MC contribution per corp
  Constants           model constants used to price tangible content
  Baseline            anchor corps and the solved baseline
  Raw deltas          corp_hand.py output (both estimands, with and without
                      hand controls)

Run corp_hand.py first. Usage: python export_corp_values.py [--baseline B]
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from corp_specs import C, CORPS, EPM, tangible

BASELINE, BASELINE_SE = 68.0, 1.2
argv = sys.argv[1:]
if "--baseline" in argv:
    BASELINE = float(argv[argv.index("--baseline") + 1])

here = Path(__file__).parent
raw = pd.read_csv(here / "corp_hand_deltas.csv")
d = raw.set_index("Corporation")

ANCHORS = {"Helion": (1.0, 1.5), "Inventrix": (2.0, 2.0), "Point Luna": (16.6, 3.0)}

# ---- main table --------------------------------------------------------------
rows = []
for name, (start, comp, ability) in CORPS.items():
    r = d.loc[name]
    obs = BASELINE + r.all_delta / EPM
    tg = tangible(name)
    rows.append({
        "Corporation": name,
        "Kept": int(r.kept), "Offered": int(r.offered),
        "Keep %": round(100 * r.kept / r.offered, 1),
        "Delta (elo, all-not-kept, hand-controlled)": round(r.all_delta, 3),
        "Delta SE": round(r.all_se, 3),
        "Delta (elo, offered-but-declined)": round(r.decl_delta, 3),
        "Observed value (MC)": round(obs, 1),
        "Value vs average corp (MC)": round(r.all_delta / EPM, 1),
        "Starting MC": start,
        "Tangible value (MC)": round(tg, 1),
        "Ability value (MC)": round(obs - tg, 1),
        "Ability SE (MC)": round(r.all_se / EPM, 1),
        "Baseline anchor": "yes" if name in ANCHORS else "",
        "Ability": ability,
    })
main = pd.DataFrame(rows).sort_values("Observed value (MC)", ascending=False)
main.to_csv(here / "corp_values.csv", index=False)

# ---- tangible breakdown ------------------------------------------------------
comp_keys = [k for k in C if any(k in comp for _, comp, _ in CORPS.values())]
brk = []
for name, (start, comp, _) in CORPS.items():
    row = {"Corporation": name, "Starting MC": start}
    for k in comp_keys:
        n = comp.get(k, 0)
        row[f"{k} (n)"] = n
        row[f"{k} (MC)"] = round(n * C[k], 1)
    row["Tangible value (MC)"] = round(tangible(name), 1)
    brk.append(row)
brk = pd.DataFrame(brk).set_index("Corporation").loc[main["Corporation"]].reset_index()

# ---- constants ---------------------------------------------------------------
const_notes = {
    "mc_prod": "MC production, per step (setup value)",
    "steel_prod": "steel production, per step", "ti_prod": "titanium production, per step",
    "energy_prod": "energy production, per step", "heat_prod": "heat production, per step",
    "plant_prod": "plant production, per step",
    "steel": "one-time steel resource (face 2)", "ti": "one-time titanium (face 3; user-set 2.5)",
    "plant": "one-time plant", "heat": "one-time heat",
    "earth": "Earth tag", "space": "Space tag", "plant_tag": "Plant tag", "jovian": "Jovian tag",
    "science": "Science tag (full-game value)",
    "building": "Building tag (only value is steel payability -> 0 on a corp)",
    "power": "Power tag", "card": "card draw at setup", "city": "city tile at setup",
    "mc": "one-time MC", "heat_prod_as_mc": "heat prod when heat is spendable as MC (Helion): MC-prod floor",
}
const = pd.DataFrame([{"Constant": k, "MC": v, "Meaning": const_notes.get(k, "")} for k, v in C.items()]
                     + [{"Constant": "EPM", "MC": EPM, "Meaning": "elo points per MC (from gen-1 card deltas)"},
                        {"Constant": "BASELINE", "MC": BASELINE, "Meaning": f"value of the average displaced corp (+/- {BASELINE_SE})"}])

# ---- baseline anchors --------------------------------------------------------
anc = []
for name, (ab, ab_se) in ANCHORS.items():
    r = d.loc[name]
    V = tangible(name) + ab
    anc.append({"Corporation": name, "Tangible (MC)": round(tangible(name), 1),
                "Assumed ability (MC)": ab, "Assumption SE": ab_se, "Total V (MC)": round(V, 1),
                "Delta": round(r.all_delta, 3),
                "Implied baseline (MC)": round(V - r.all_delta / EPM, 2),
                "SE": round(((r.all_se / EPM) ** 2 + ab_se ** 2) ** 0.5, 2)})
anc = pd.DataFrame(anc)
w = 1 / anc["SE"] ** 2
anc.loc[len(anc)] = {"Corporation": "Weighted baseline", "Implied baseline (MC)": round((anc["Implied baseline (MC)"] * w).sum() / w.sum(), 2),
                     "SE": round(w.sum() ** -0.5, 2)}

# ---- write workbook ----------------------------------------------------------
wb = Workbook()
wb.remove(wb.active)


def sheet(title, df, wrap_cols=()):
    ws = wb.create_sheet(title)
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    for i, col in enumerate(df.columns, 1):
        width = 60 if col in wrap_cols else min(max(len(str(col)) * 0.6 + 4,
                                                 max((len(str(v)) for v in df[col]), default=0) + 2), 34)
        ws.column_dimensions[get_column_letter(i)].width = width
        if col in wrap_cols:
            for cell in ws[get_column_letter(i)][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"


sheet("Corp values", main, wrap_cols=("Ability",))
sheet("Tangible breakdown", brk)
sheet("Constants", const, wrap_cols=("Meaning",))
sheet("Baseline", anc)
sheet("Raw deltas", raw)
out = here / "corp_values.xlsx"
wb.save(out)
print(f"Wrote {out} and {here / 'corp_values.csv'}")
print(main[["Corporation", "Observed value (MC)", "Tangible value (MC)", "Ability value (MC)"]].to_string(index=False))
